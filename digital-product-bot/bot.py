from __future__ import annotations

import csv
import json
import math
import os
import re
import shutil
import sys
import time
import zipfile
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Tuple
from urllib.parse import quote_plus

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
CONFIG_PATH = BASE_DIR / "config.json"
USER_AGENT = "RCS-Digital-Product-Scout/1.0 (product research bot; respectful rate limits)"


@dataclass
class Opportunity:
    phrase: str
    score: int
    signal_count: int
    sources: List[str]
    kind: str
    reason: str


def load_config() -> dict:
    with CONFIG_PATH.open("r", encoding="utf-8") as f:
        config = json.load(f)

    craft_seed_path = BASE_DIR / "craft_factory" / "craft_market_seeds.json"
    if craft_seed_path.exists():
        try:
            craft_payload = json.loads(craft_seed_path.read_text(encoding="utf-8"))
            extra_seeds = []
            for values in craft_payload.values():
                if isinstance(values, list):
                    extra_seeds.extend(str(v) for v in values)
            config["seed_markets"] = list(dict.fromkeys(config.get("seed_markets", []) + extra_seeds))
        except Exception:
            # Discovery should continue even if optional craft seed configuration is malformed.
            pass

    return config


def clean_phrase(text: str) -> str:
    text = re.sub(r"https?://\S+", "", text.lower())
    text = re.sub(r"[^a-z0-9%+&' /_-]", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" -_/.")
    return text[:140]


def slugify(text: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return value[:70] or "product"


def google_suggestions(session: requests.Session, query: str, limit: int) -> List[str]:
    url = "https://suggestqueries.google.com/complete/search"
    params = {"client": "firefox", "q": query, "hl": "en"}
    response = session.get(url, params=params, timeout=12)
    response.raise_for_status()
    payload = response.json()
    if not isinstance(payload, list) or len(payload) < 2 or not isinstance(payload[1], list):
        return []
    return [clean_phrase(x) for x in payload[1][:limit] if clean_phrase(x)]


def reddit_titles(session: requests.Session, community: str, limit: int) -> List[str]:
    url = f"https://www.reddit.com/r/{quote_plus(community)}/search.json"
    params = {
        "q": "template OR spreadsheet OR tracker OR planner OR checklist OR printable OR papercraft OR svg OR pattern OR favor box",
        "restrict_sr": "on",
        "sort": "new",
        "t": "month",
        "limit": min(limit, 50),
        "raw_json": 1,
    }
    response = session.get(url, params=params, timeout=12)
    response.raise_for_status()
    children = response.json().get("data", {}).get("children", [])
    titles = []
    for child in children:
        title = clean_phrase(child.get("data", {}).get("title", ""))
        if title:
            titles.append(title)
    return titles


def discover_signals(config: dict) -> Tuple[Dict[str, Counter], List[str]]:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "en-US,en;q=0.8"})
    signals: Dict[str, Counter] = defaultdict(Counter)
    errors: List[str] = []

    for seed in config["seed_markets"]:
        try:
            for suggestion in google_suggestions(
                session, seed, int(config.get("google_suggest_limit_per_seed", 8))
            ):
                signals[suggestion]["google_suggest"] += 1
            time.sleep(0.15)
        except Exception as exc:
            errors.append(f"Google Suggest failed for '{seed}': {exc}")

    for community in config.get("reddit_communities", []):
        try:
            for title in reddit_titles(
                session, community, int(config.get("reddit_posts_per_query", 20))
            ):
                signals[title][f"reddit:{community}"] += 1
            time.sleep(0.35)
        except Exception as exc:
            errors.append(f"Reddit failed for r/{community}: {exc}")

    # Ensure a graceful, useful first run even when an external source blocks requests.
    if not signals:
        for seed in config["seed_markets"]:
            signals[clean_phrase(seed)]["configured_seed_fallback"] += 1

    return signals, errors


def infer_kind(phrase: str) -> str:
    p = phrase.lower()
    if any(term in p for term in ["favor box", "favour box", "treat box", "party box", "cube box", "papercraft"]):
        return "papercraft"
    if any(term in p for term in ["svg", "cut file", "cricut", "silhouette", "laser cut", "dxf"]):
        return "cut-file"
    if any(term in p for term in ["leather pattern", "sewing pattern", "woodworking plan", "paper model"]):
        return "craft-pattern"
    if "checklist" in p or "onboarding" in p or "maintenance" in p:
        return "checklist"
    if "planner" in p or "calendar" in p or "schedule" in p:
        return "planner"
    if "calculator" in p or "pricing" in p or "profit" in p or "margin" in p:
        return "calculator"
    if "tracker" in p or "spreadsheet" in p or "excel" in p or "sheets" in p:
        return "tracker"
    return "toolkit"


def score_phrase(phrase: str, counter: Counter, config: dict) -> Opportunity | None:
    p = phrase.lower()
    excluded = [term.lower() for term in config.get("excluded_terms", [])]
    if any(term in p for term in excluded):
        return None

    non_product_intent = [term.lower() for term in config.get("non_product_intent_terms", [])]
    if any(term in p for term in non_product_intent):
        return None

    intent_terms = [term.lower() for term in config.get("commercial_intent_terms", [])]
    intent_hits = sum(1 for term in intent_terms if term in p)
    words = [w for w in re.split(r"\s+", p) if w]
    source_count = len(counter)
    signal_count = sum(counter.values())

    demand = min(35, 12 + (signal_count - 1) * 7 + (source_count - 1) * 8)
    intent = min(25, intent_hits * 12)
    specificity = min(15, max(0, (len(words) - 2) * 4))
    buildable_terms = [
        "template", "spreadsheet", "tracker", "planner", "checklist", "calculator",
        "dashboard", "calendar", "worksheet", "onboarding", "budget", "inventory",
        "project", "invoice", "quote", "content", "maintenance", "fitness",
        "printable", "papercraft", "favor box", "favour box", "treat box", "party box",
        "cube box", "svg", "cut file", "cricut", "silhouette", "pattern",
        "leather", "sewing", "woodworking", "laser cut", "paper model", "party kit"
    ]
    buildable_hits = sum(1 for term in buildable_terms if term in p)
    buildability = min(24, 8 + buildable_hits * 4)

    craft_terms = [
        "papercraft", "favor box", "favour box", "treat box", "party box", "cube box",
        "printable", "svg", "cut file", "cricut", "silhouette", "leather pattern",
        "sewing pattern", "woodworking plan", "laser cut", "paper model"
    ]
    craft_bonus = 8 if any(term in p for term in craft_terms) else 0
    generic_penalty = 10 if len(words) <= 2 else 0
    question_penalty = 8 if p.startswith(("why ", "what ", "how ", "can ", "should ")) else 0

    raw = demand + intent + specificity + buildability + craft_bonus - generic_penalty - question_penalty
    score = max(0, min(100, int(round(raw))))
    kind = infer_kind(p)
    reason = (
        f"demand={demand}, intent={intent}, specificity={specificity}, "
        f"buildability={buildability}, craft_bonus={craft_bonus}, penalties={generic_penalty + question_penalty}"
    )
    return Opportunity(
        phrase=phrase,
        score=score,
        signal_count=signal_count,
        sources=sorted(counter.keys()),
        kind=kind,
        reason=reason,
    )


def build_opportunities(signals: Dict[str, Counter], config: dict) -> List[Opportunity]:
    opportunities: List[Opportunity] = []
    for phrase, counter in signals.items():
        item = score_phrase(phrase, counter, config)
        if item:
            opportunities.append(item)

    # De-duplicate close phrases by slug-like canonical form.
    best_by_key: Dict[str, Opportunity] = {}
    for item in opportunities:
        key_words = [
            w for w in item.phrase.lower().split()
            if w not in {"best", "free", "download", "printable", "the", "a", "for", "to"}
        ]
        key = " ".join(key_words[:8])
        previous = best_by_key.get(key)
        if previous is None or item.score > previous.score:
            best_by_key[key] = item

    return sorted(best_by_key.values(), key=lambda x: (x.score, x.signal_count), reverse=True)


def headers_for_phrase(phrase: str) -> List[str]:
    p = phrase.lower()
    if "budget" in p or "expense" in p or "cash flow" in p:
        return ["Date", "Category", "Description", "Budget", "Actual", "Variance", "Notes"]
    if "content" in p or "social media" in p:
        return ["Publish Date", "Channel", "Topic", "Format", "Status", "CTA", "Notes"]
    if "project" in p or "task" in p:
        return ["Task", "Owner", "Priority", "Status", "Start Date", "Due Date", "Progress %", "Notes"]
    if "inventory" in p or "stock" in p:
        return ["SKU", "Item", "Category", "Supplier", "Qty On Hand", "Reorder Level", "Unit Cost", "Stock Value"]
    if "invoice" in p or "quote" in p:
        return ["Reference", "Client", "Issue Date", "Due Date", "Amount", "Status", "Paid Date", "Notes"]
    if "fitness" in p or "workout" in p or "training" in p:
        return ["Date", "Session", "Duration", "Distance", "Intensity", "RPE", "Status", "Notes"]
    if "property" in p or "maintenance" in p:
        return ["Area", "Task", "Frequency", "Last Done", "Next Due", "Responsible", "Status", "Notes"]
    if "onboarding" in p:
        return ["Step", "Owner", "Due Timing", "Status", "Required Info", "Completed Date", "Notes"]
    return ["Item", "Category", "Owner", "Priority", "Status", "Due Date", "Value", "Notes"]


def autosize(ws) -> None:
    for column_cells in ws.columns:
        width = 10
        for cell in column_cells:
            if cell.value is not None:
                width = max(width, min(36, len(str(cell.value)) + 2))
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def create_workbook(path: Path, title: str, phrase: str, headers: List[str]) -> None:
    wb = Workbook()
    start = wb.active
    start.title = "Start Here"
    start["A1"] = title
    start["A1"].font = Font(size=18, bold=True)
    start["A3"] = "Purpose"
    start["B3"] = f"A practical, editable tool built around the demand phrase: {phrase}."
    start["A5"] = "How to use"
    start["B5"] = "Add your own records on the Main sheet. Use filters to review status, dates, owners and values."
    start["A7"] = "Commercial status"
    start["B7"] = "Draft product - review required before sale or publication."
    start.column_dimensions["A"].width = 22
    start.column_dimensions["B"].width = 80
    start["B3"].alignment = Alignment(wrap_text=True)
    start["B5"].alignment = Alignment(wrap_text=True)

    main = wb.create_sheet("Main")
    for col, header in enumerate(headers, start=1):
        c = main.cell(row=1, column=col, value=header)
        c.font = Font(bold=True)
    main.freeze_panes = "A2"
    main.auto_filter.ref = f"A1:{get_column_letter(len(headers))}51"

    for row in range(2, 52):
        for col in range(1, len(headers) + 1):
            main.cell(row=row, column=col, value="")

    if "Variance" in headers:
        b = headers.index("Budget") + 1
        a = headers.index("Actual") + 1
        v = headers.index("Variance") + 1
        for row in range(2, 52):
            main.cell(row=row, column=v, value=f"={get_column_letter(b)}{row}-{get_column_letter(a)}{row}")
    if "Stock Value" in headers:
        q = headers.index("Qty On Hand") + 1
        c = headers.index("Unit Cost") + 1
        s = headers.index("Stock Value") + 1
        for row in range(2, 52):
            main.cell(row=row, column=s, value=f"={get_column_letter(q)}{row}*{get_column_letter(c)}{row}")

    autosize(main)

    dash = wb.create_sheet("Dashboard")
    dash["A1"] = "Quick Dashboard"
    dash["A1"].font = Font(size=16, bold=True)
    dash["A3"] = "Rows available"
    dash["B3"] = 50
    dash["A4"] = "Records entered"
    dash["B4"] = "=COUNTA(Main!A2:A51)"
    if "Status" in headers:
        status_col = get_column_letter(headers.index("Status") + 1)
        dash["A6"] = "Completed / Done"
        dash["B6"] = f'=COUNTIF(Main!{status_col}2:{status_col}51,"*done*")+COUNTIF(Main!{status_col}2:{status_col}51,"*complete*")'
    autosize(dash)
    wb.save(path)


def title_case_phrase(phrase: str) -> str:
    small_words = {"a", "an", "and", "for", "of", "the", "to", "with"}
    words = phrase.split()
    result = []
    for i, word in enumerate(words):
        result.append(word.lower() if i and word.lower() in small_words else word.capitalize())
    return " ".join(result)


def create_guide(path: Path, product_title: str, phrase: str, headers: List[str]) -> None:
    header_list = "\n".join(f"- {h}" for h in headers)
    text = f"""# {product_title}\n\n## What this product does\n\nThis editable toolkit is designed around the problem represented by **{phrase}**. It gives the buyer a simple working system rather than a blank document.\n\n## Included fields\n\n{header_list}\n\n## Suggested workflow\n\n1. Duplicate the original file before customising it.\n2. Add existing records to the Main sheet.\n3. Standardise status and category names so filtering works cleanly.\n4. Review the Dashboard sheet during a weekly check-in.\n5. Archive completed items periodically instead of deleting useful history.\n\n## Quality checklist before first use\n\n- [ ] Replace sample wording with your own terminology.\n- [ ] Confirm date and currency formats for your region.\n- [ ] Add or remove columns that do not fit your workflow.\n- [ ] Test formulas after structural changes.\n- [ ] Save a clean master copy.\n\n## Important\n\nThis is a general productivity template. It is not legal, tax, financial, medical, or other regulated professional advice.\n"""
    path.write_text(text, encoding="utf-8")


def create_csv(path: Path, headers: List[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for _ in range(5):
            writer.writerow([""] * len(headers))


def keyword_list(phrase: str, kind: str) -> List[str]:
    base = [phrase, f"{kind} template", "digital download", "editable spreadsheet", "productivity tool"]
    words = [w for w in phrase.lower().split() if len(w) > 3]
    base.extend([f"{w} template" for w in words[:4]])
    seen = []
    for item in base:
        item = clean_phrase(item)
        if item and item not in seen:
            seen.append(item)
    return seen[:12]


def suggested_price(kind: str, config: dict, score: int) -> dict:
    low, high = config.get("price_bands_usd", {}).get(kind, [9, 25])
    midpoint = low + (high - low) * min(1.0, max(0.0, (score - 50) / 50))
    usd = round(midpoint, 2)
    # Approximate display-only ZAR band; seller should set live marketplace currency at listing time.
    zar_reference = round(usd * 18.0)
    return {"usd": usd, "zar_reference_at_18_per_usd": zar_reference, "band_usd": [low, high]}


def create_product(opportunity: Opportunity, product_root: Path, config: dict) -> Path:
    slug = slugify(opportunity.phrase)
    folder = product_root / slug
    folder.mkdir(parents=True, exist_ok=True)

    product_title = f"{title_case_phrase(opportunity.phrase)} - Editable {opportunity.kind.capitalize()}"
    headers = headers_for_phrase(opportunity.phrase)

    create_workbook(folder / f"{slug}.xlsx", product_title, opportunity.phrase, headers)
    create_csv(folder / f"{slug}-starter.csv", headers)
    create_guide(folder / "buyer-guide.md", product_title, opportunity.phrase, headers)

    listing = {
        "title": product_title[:140],
        "short_description": (
            f"A ready-to-edit {opportunity.kind} for people searching for '{opportunity.phrase}'. "
            "Includes an Excel workbook, starter CSV and practical setup guide."
        ),
        "long_description": (
            f"Turn {opportunity.phrase} into a repeatable workflow. This digital download includes an editable "
            "Excel workbook with a Start Here sheet, working table and quick dashboard, plus a CSV starter file "
            "and concise buyer guide. Designed for practical everyday use and easy customisation."
        ),
        "keywords": keyword_list(opportunity.phrase, opportunity.kind),
        "suggested_price": suggested_price(opportunity.kind, config, opportunity.score),
        "source_phrase": opportunity.phrase,
        "opportunity_score": opportunity.score,
        "source_signals": opportunity.sources,
        "disclaimer": "General productivity template only; not professional legal, tax, financial or medical advice.",
    }
    (folder / "listing.json").write_text(json.dumps(listing, indent=2), encoding="utf-8")

    status = {
        "status": "AWAITING_APPROVAL",
        "approved_by": None,
        "approved_at": None,
        "publishing_enabled": False,
    }
    (folder / "status.json").write_text(json.dumps(status, indent=2), encoding="utf-8")

    manifest = {
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "opportunity": asdict(opportunity),
        "files": sorted(p.name for p in folder.iterdir() if p.is_file()),
    }
    (folder / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    zip_path = product_root / f"{slug}.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for item in folder.iterdir():
            if item.is_file():
                zf.write(item, arcname=f"{slug}/{item.name}")
    return zip_path


def main() -> int:
    config = load_config()
    run_date = datetime.now(timezone.utc).date().isoformat()
    output_base = Path(os.environ.get("OUTPUT_DIR", str(BASE_DIR / "output")))
    run_root = output_base / run_date
    products_root = run_root / "products"
    products_root.mkdir(parents=True, exist_ok=True)

    signals, errors = discover_signals(config)
    opportunities = build_opportunities(signals, config)
    minimum = int(config.get("minimum_score", 62))
    selected = [o for o in opportunities if o.score >= minimum]

    # If live evidence is thin, still generate the highest ranked draft so the run remains inspectable.
    if not selected and opportunities:
        selected = opportunities[:1]

    selected = selected[: int(config.get("max_products_per_run", 3))]
    packages = []
    for opportunity in selected:
        packages.append(str(create_product(opportunity, products_root, config).relative_to(run_root)))

    report = {
        "run_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": "COMPLETE",
        "publishing_performed": False,
        "signals_seen": len(signals),
        "opportunities_scored": len(opportunities),
        "minimum_score": minimum,
        "selected": [asdict(o) for o in selected],
        "top_20": [asdict(o) for o in opportunities[:20]],
        "packages": packages,
        "source_errors": errors,
    }
    (run_root / "run_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")

    print(json.dumps(report, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
